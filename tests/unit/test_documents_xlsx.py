"""Excel to Markdown: one heading per sheet, a pipe table of stored values, capped rows."""

import io
import re
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest

from mcp_connector.documents import limits
from mcp_connector.documents import xlsx as xlsx_conv
from mcp_connector.errors import REASON_GUARD_TRIPPED, ToolError

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "documents" / "sample.xlsx"


def test_each_sheet_becomes_a_heading_and_a_pipe_table() -> None:
    text = xlsx_conv.to_markdown(FIXTURE.read_bytes())

    assert "## Report" in text
    assert "## Appendix" in text
    assert "| Item | Amount |" in text
    assert "| Servers | 12 |" in text
    assert "| Date | 2026-09-30 |" in text
    assert text.index("## Report") < text.index("## Appendix")


def test_trailing_empty_rows_and_columns_are_trimmed() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "a"
    sheet["B1"] = "b"
    sheet["A2"] = "c"
    sheet["E9"] = None
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = xlsx_conv.to_markdown(buffer.getvalue())
    lines = [line for line in text.splitlines() if line.startswith("|")]
    assert lines == ["| a | b |", "| --- | --- |", "| c |  |"]


def test_rows_above_the_cap_are_cut_with_a_note(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(xlsx_conv, "MAX_ROWS", 3)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for i in range(10):
        sheet.append([i])
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = xlsx_conv.to_markdown(buffer.getvalue())
    assert text.count("\n| ") <= 4  # header row, separator, two more rows at most
    assert "(more than 3 rows; the rest is omitted)" in text


def test_the_row_note_stays_when_the_kept_rows_are_all_empty(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(xlsx_conv, "MAX_ROWS", 3)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A10"] = "late"
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = xlsx_conv.to_markdown(buffer.getvalue())
    assert "(more than 3 rows; the rest is omitted)" in text


def _with_dimension(data: bytes, ref: str) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                payload = re.sub(
                    rb'<dimension ref="[^"]*"/>', f'<dimension ref="{ref}"/>'.encode(), payload
                )
            target.writestr(info.filename, payload)
    return out.getvalue()


def test_a_lying_dimension_does_not_pad_every_row_to_16384_columns(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for i in range(2000):
        sheet.append([i])
    buffer = io.BytesIO()
    workbook.save(buffer)
    data = _with_dimension(buffer.getvalue(), "A1:XFD2000")

    # Every cell the converter sees passes through _text. Padding to the declared 16384
    # columns would be 32 million cells; the cap allows 256 per row at most.
    bound = 2000 * 256
    seen = 0
    real_text = xlsx_conv._text

    def counting(value: object) -> str:
        nonlocal seen
        seen += 1
        if seen > bound:
            raise AssertionError("the converter padded rows beyond the column cap")
        return real_text(value)

    monkeypatch.setattr(xlsx_conv, "_text", counting)
    started = time.perf_counter()
    text = xlsx_conv.to_markdown(data)
    assert time.perf_counter() - started < 5
    assert "| 1999 |" in text


def test_columns_above_the_cap_are_not_read() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.cell(row=1, column=1, value="first")
    sheet.cell(row=1, column=xlsx_conv.MAX_COLS + 1, value="beyond")
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = xlsx_conv.to_markdown(buffer.getvalue())
    assert "beyond" not in text
    rows = [line for line in text.splitlines() if line.startswith("|")]
    assert all(row.count(" | ") + 1 <= xlsx_conv.MAX_COLS for row in rows)


def test_too_many_sheets_are_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(xlsx_conv, "MAX_SHEETS", 2)
    workbook = openpyxl.Workbook()
    workbook.create_sheet("Two")
    workbook.create_sheet("Three")
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(ToolError) as info:
        xlsx_conv.to_markdown(buffer.getvalue())
    assert info.value.reason == REASON_GUARD_TRIPPED


def test_the_output_stops_at_the_cap_with_a_note(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(limits, "MAX_OUTPUT_CHARS", 40)
    text = xlsx_conv.to_markdown(FIXTURE.read_bytes())
    body, _, note = text.rstrip("\n").rpartition("\n")
    assert note == "(output truncated at 40 characters)"
    assert len(body.rstrip("\n")) <= 40
    assert "## Appendix" not in text
