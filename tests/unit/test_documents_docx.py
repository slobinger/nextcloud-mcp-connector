"""Word to Markdown: headings by level, paragraphs, list items, pipe tables, dropped rest."""

import copy
import io
import time
from pathlib import Path

import docx
import docx.document
import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from mcp_connector.documents import docx as docx_conv
from mcp_connector.documents import limits

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "documents" / "sample.docx"


def test_headings_paragraphs_lists_and_tables_come_out_in_document_order() -> None:
    text = docx_conv.to_markdown(FIXTURE.read_bytes())
    lines = [line for line in text.splitlines() if line.strip()]

    assert lines[0] == "# Quarterly Report"
    assert lines[1] == "The figures below are final."
    assert lines[2] == "- First point"
    assert lines[3] == "| Item | Amount |"
    assert lines[4] == "| --- | --- |"
    assert lines[5] == "| Servers | 12 |"
    assert lines[6] == "## Appendix"
    assert lines[7] == "Nothing else."


def test_heading_levels_above_six_are_capped() -> None:
    assert docx_conv._heading_prefix("Heading 9") == "######"
    assert docx_conv._heading_prefix("Heading 2") == "##"
    assert docx_conv._heading_prefix("Title") == "#"
    assert docx_conv._heading_prefix("Normal") == ""


def test_cell_text_is_flattened_to_one_line_and_pipes_are_escaped() -> None:
    assert docx_conv._cell("a\nb | c") == "a b \\| c"


def _saved(document: docx.document.Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_a_huge_grid_span_does_not_repeat_the_cell() -> None:
    document = docx.Document()
    table = document.add_table(rows=50, cols=1)
    for row in table.rows:
        row.cells[0].text = "v"
        span = OxmlElement("w:gridSpan")
        span.set(qn("w:val"), "100000")
        row._tr.tc_lst[0].get_or_add_tcPr().append(span)

    started = time.perf_counter()
    text = docx_conv.to_markdown(_saved(document))
    assert time.perf_counter() - started < 5
    rows = [line for line in text.splitlines() if line.startswith("| v")]
    assert len(rows) == 50
    assert all(row == "| v |" for row in rows)


def test_a_merged_cell_appears_once() -> None:
    document = docx.Document()
    table = document.add_table(rows=1, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "wide"
    table.cell(0, 2).text = "narrow"

    text = docx_conv.to_markdown(_saved(document))
    assert text.splitlines()[0] == "| wide | narrow |"


def test_table_columns_and_rows_are_capped(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(docx_conv, "MAX_TABLE_COLS", 2)
    monkeypatch.setattr(docx_conv, "MAX_TABLE_ROWS", 3)
    document = docx.Document()
    table = document.add_table(rows=5, cols=4)
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"r{r}c{c}"

    text = docx_conv.to_markdown(_saved(document))
    assert "| r0c0 | r0c1 |" in text
    assert "r0c2" not in text
    assert "r3c0" not in text
    assert "(more than 3 rows; the rest of the table is omitted)" in text
    assert "(more than 2 columns; the rest of each row is omitted)" in text


def test_twenty_thousand_paragraphs_convert_quickly() -> None:
    document = docx.Document()
    document.add_heading("Top", level=1)
    paragraph = document.add_paragraph("line")._p
    body = paragraph.getparent()
    assert body is not None
    for _ in range(19999):
        body.append(copy.deepcopy(paragraph))

    started = time.perf_counter()
    text = docx_conv.to_markdown(_saved(document))
    elapsed = time.perf_counter() - started
    assert text.startswith("# Top\n")
    assert text.count("line") == 20000
    assert elapsed < 2


def test_the_output_stops_at_the_cap_with_a_note(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(limits, "MAX_OUTPUT_CHARS", 40)
    text = docx_conv.to_markdown(FIXTURE.read_bytes())
    body, _, note = text.rstrip("\n").rpartition("\n")
    assert note == "(output truncated at 40 characters)"
    assert len(body.rstrip("\n")) <= 40
    assert text.startswith("# Quarterly Report")
