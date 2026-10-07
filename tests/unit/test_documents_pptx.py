"""PowerPoint to Markdown: a heading per slide, text frames, tables, speaker notes."""

import io
from pathlib import Path

import pptx
import pytest
from pptx.util import Inches

from mcp_connector.documents import limits
from mcp_connector.documents import pptx as pptx_conv

FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "documents" / "sample.pptx"


def test_each_slide_becomes_a_heading_with_its_text_tables_and_notes() -> None:
    text = pptx_conv.to_markdown(FIXTURE.read_bytes())

    assert "## Slide 1: Quarterly Report" in text
    assert "The figures below are final." in text
    assert "| Item | Amount |" in text
    assert "| Servers | 12 |" in text
    assert "Notes: Speaker note one." in text
    assert "## Slide 2: Appendix" in text
    assert text.index("## Slide 1") < text.index("## Slide 2")


def test_the_title_is_not_repeated_as_a_text_frame() -> None:
    text = pptx_conv.to_markdown(FIXTURE.read_bytes())
    assert text.count("Quarterly Report") == 1


def test_the_output_stops_at_the_cap_with_a_note(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(limits, "MAX_OUTPUT_CHARS", 40)
    text = pptx_conv.to_markdown(FIXTURE.read_bytes())
    body, _, note = text.rstrip("\n").rpartition("\n")
    assert note == "(output truncated at 40 characters)"
    assert len(body.rstrip("\n")) <= 40
    assert "## Slide 2" not in text


def test_a_soft_line_break_gives_two_lines() -> None:
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.add_run().text = "Line one"
    paragraph.add_line_break()
    paragraph.add_run().text = "Line two"
    buffer = io.BytesIO()
    presentation.save(buffer)

    lines = pptx_conv.to_markdown(buffer.getvalue()).splitlines()
    assert "Line one" in lines
    assert "Line two" in lines
    assert not any("Line oneLine two" in line for line in lines)
