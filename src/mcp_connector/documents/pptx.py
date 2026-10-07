"""PowerPoint to Markdown: one section per slide, in slide order. Images are dropped."""

import io

from .docx import _cell, _table_lines
from .limits import Output

__all__ = ["to_markdown"]


def to_markdown(data: bytes) -> str:
    import pptx
    from pptx.shapes.autoshape import Shape
    from pptx.shapes.graphfrm import GraphicFrame

    presentation = pptx.Presentation(io.BytesIO(data))
    out = Output()
    for number, slide in enumerate(presentation.slides, start=1):
        if out.full:
            break
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        out.add(f"## Slide {number}: {title}" if title else f"## Slide {number}")
        out.add("")
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if isinstance(shape, Shape) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    # paragraph.text writes a soft line break as a vertical tab, which
                    # splitlines treats as a line boundary.
                    for line in paragraph.text.splitlines():
                        if line.strip():
                            out.add(line.strip())
                if shape.text_frame.text.strip():
                    out.add("")
            elif isinstance(shape, GraphicFrame) and shape.has_table:
                rows = [[_cell(cell.text) for cell in row.cells] for row in shape.table.rows]
                out.extend(_table_lines(rows))
                out.add("")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                out.add(f"Notes: {' '.join(notes.text.split())}")
                out.add("")
    return out.text()
